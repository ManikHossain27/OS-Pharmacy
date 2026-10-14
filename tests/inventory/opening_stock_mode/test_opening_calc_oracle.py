"""Opening Stock mode - calculation model (no browser).

Verifies ``utils/opening_calc.py`` (BTP + VAT/Box derived from Retail x (1 - Opening Stock %), then the Default-mode
calculation) against the spec worked examples, the numbers in the user's workbook (sheet ``Opening Stock``) and
invariants over the full combination matrix. ``tools/validate_opening_oracle.py`` additionally proves the oracle
against the recalculated workbook cell by cell (8,470 comparisons, 0 mismatches).

The last class pins the structural rules of the (corrected 2026-10-12) workbook the oracle relies on.
"""
from __future__ import annotations

import random
from pathlib import Path

import allure
import openpyxl
import pytest

from utils.inventory_calc import Line, build as build_default
from utils.opening_calc import (DEFAULT_PCT, btp_exact, btp_for, build, clamp_pct, excel_round, opening_line,
                                pct_for, sheet_due_advance)
from utils.opening_scenarios import calc_matrix

pytestmark = [pytest.mark.inventory, pytest.mark.calc]

TOL = 1e-6
WORKBOOK = Path(__file__).resolve().parents[3] / "Knowledge/Inventory Management/PMS_Innventory Calculation (Non Rounded).xlsx"


def _fails(items, check):
    out = []
    for item in items:
        msg = check(item)
        if msg:
            out.append(f"{item[0]}: {msg}")
    return out


@allure.epic("Inventory")
@allure.feature("Opening Stock mode - calculation model")
class TestBtpDerivation:
    """BTP + VAT/Box = Retail Price/Box x (1 - Opening Stock % / 100) - spec sections 15-19, sheet col I."""

    @allure.story("Spec 15: the default Opening Stock % is 12")
    def test_default_percentage(self):
        assert DEFAULT_PCT == 12

    @allure.story("Spec 18: retail 100 at the default 12 % -> BTP + VAT/Box 88")
    def test_default_derivation(self):
        assert btp_for(100) == 88 and btp_for(200) == 176 and btp_for(150) == 132

    @allure.story("Spec 19 worked example: retail 100, BTP + VAT/Box 80, Box Qty 10 -> Total 800, Opening Stock % = 20")
    def test_manual_override_example(self):
        line = opening_line(10, 10, 100, btp=80)
        assert line.total_btp_vat == 800
        assert pct_for(100, 80) == pytest.approx(20)

    @allure.story("Spec 19: the % implied by a manual BTP + VAT/Box (retail - btp) / retail")
    @pytest.mark.parametrize("retail,btp,pct", [(100, 88, 12), (200, 150, 25), (50, 50, 0), (40, 0, 100), (100, 120, -20)])
    def test_pct_for(self, retail, btp, pct):
        assert pct_for(retail, btp) == pytest.approx(pct)

    @allure.story("The application rounds BTP + VAT/Box to 2 decimals (verified live: 33.33 at 12 % -> 29.33)")
    @pytest.mark.parametrize("retail,pct,expected", [(33.33, 12, 29.33), (33.33, 7.5, 30.83), (150, 12.345, 131.48),
                                                      (250, 12.345, 219.14), (0.01, 12, 0.01), (999.99, 33.333, 666.66)])
    def test_two_decimal_rounding(self, retail, pct, expected):
        assert btp_for(retail, pct) == pytest.approx(expected, abs=1e-9)

    @allure.story("0 % -> BTP = Retail; 100 % -> BTP = 0")
    def test_bounds(self):
        assert btp_for(123.45, 0) == 123.45 and btp_for(123.45, 100) == 0

    @allure.story("The % is clamped into 0..100 by the form (verified live: -5 -> 0, 150 -> 100)")
    @pytest.mark.parametrize("typed,clamped", [(-5, 0), (0, 0), (12, 12), (100, 100), (150, 100), (1000, 100)])
    def test_clamp(self, typed, clamped):
        assert clamp_pct(typed) == clamped
        assert btp_for(100, typed) == btp_for(100, clamped)

    @allure.story("BTP + VAT/Box falls as the % rises and rises with the retail price")
    def test_monotonic(self):
        by_pct = [btp_for(100, p) for p in range(0, 101, 5)]
        assert by_pct == sorted(by_pct, reverse=True)
        by_retail = [btp_for(r, 12) for r in (10, 20, 50, 100, 500)]
        assert by_retail == sorted(by_retail)

    @allure.story("The sheet keeps the exact product; the app's 2-dp BTP differs by at most 0.005 per box")
    def test_app_vs_sheet_rounding_difference_is_tiny(self):
        rng = random.Random(11)
        worst = 0.0
        for _ in range(2000):
            r, p = round(rng.uniform(0.01, 1500), 2), rng.choice([0, 7.5, 12, 12.345, 20, 33.333, 99.9])
            worst = max(worst, abs(btp_for(r, p) - btp_exact(r, p)))
        assert worst <= 0.005 + 1e-9, worst


@allure.epic("Inventory")
@allure.feature("Opening Stock mode - calculation model")
class TestWorkbookAnchors:
    """Numbers lifted straight from the workbook sheet ``Opening Stock`` (cached cell values)."""

    @allure.story("Sheet rows 8-10: 3 Bion (free 2, BTP 700) + special 20 on a box-0 product -> Final 3480")
    def test_sheet_rows(self):
        inv = build([Line(5, 60, 700, free_qty=2, retail_per_box=720, wholesale_per_box=700),
                     Line(0, 1, 171.6, retail_per_box=195, wholesale_per_box=195, special_discount=20),
                     Line(0, 12, 369.6, retail_per_box=420, wholesale_per_box=420)],
                    less_basis="purchased")
        l = inv.lines[0]
        assert (inv.total_price, inv.total_discount, inv.final_price) == (3500, 20, 3480)
        assert l.purchase_price_pc == pytest.approx(11.589403973509933, abs=TOL)
        assert l.purchase_price_box == pytest.approx(695.3642384105959, abs=TOL)
        assert l.retail_margin_pct == pytest.approx(0.034216335540838985, abs=TOL)
        assert l.retail_margin_value == pytest.approx(24.63576158940407, abs=TOL)
        assert l.wholesale_margin_pct == pytest.approx(0.006622516556291527, abs=TOL)
        assert l.wholesale_margin_value == pytest.approx(4.635761589404069, abs=TOL)

    @allure.story("Sheet N2:N4 - supplier balance 6337 + Final 3480 - paid 0 = Due 9817, Advance 0")
    def test_sheet_supplier_due(self):
        inv = build([Line(5, 60, 700, free_qty=2), Line(0, 1, 171.6, special_discount=20)], supplier_balance=6337)
        assert inv.final_price == 3480 and inv.due == 9817 and inv.advance == 0

    @allure.story("Opening Stock and Default calculate identically once BTP + VAT/Box is known (same engine)")
    def test_same_engine_as_default(self):
        lines = [opening_line(2, 10, 100), opening_line(3, 12, 200, pct=15, free_qty=5), opening_line(1, 1, 120, pct=20)]
        a = build(lines, pct=10, paid=100, supplier_balance=50)
        b = build_default([Line(l.box_qty, l.qty_per_box, l.btp_vat_per_box, free_qty=l.free_qty,
                                retail_per_box=l.retail_per_box) for l in lines], pct=10, paid=100, supplier_balance=50)
        assert (a.final_price, a.total_discount, a.due) == (b.final_price, b.total_discount, b.due)
        assert [l.purchase_price_pc for l in a.lines] == [l.purchase_price_pc for l in b.lines]


@allure.epic("Inventory")
@allure.feature("Opening Stock mode - calculation model")
class TestCombinationMatrix:
    """Every (box x qty/box x retail x opening % x free x discount) combination of ``calc_matrix``."""

    @pytest.fixture(scope="class")
    def solved(self):
        return [(cid, build(lines, pct=pct, flat=flat), pct, flat) for cid, lines, pct, flat in calc_matrix()]

    @allure.story("Matrix is large enough to be meaningful")
    def test_matrix_size(self, solved):
        assert len(solved) > 3_000

    @allure.story("Total BTP + VAT = sum of (derived BTP + VAT/Box x Box Qty)")
    def test_total_is_sum_of_lines(self, solved):
        assert not _fails(solved, lambda s: None if abs(s[1].total_price - sum(l.btp_vat_per_box * l.box_qty for l in s[1].lines)) < TOL
                          else "total != sum")

    @allure.story("Free-qty lines never receive a discount; the percentage applies to every eligible line")
    def test_free_lines_zero_discount(self, solved):
        assert not _fails(solved, lambda s: next((f"free line discounted {l.discount_value}" for l in s[1].lines
                                                  if l.free_qty > 0 and l.discount_value != 0), None))

    @allure.story("Percentage discount = pct x eligible base; Flat discount = flat (or 0 when nothing is eligible)")
    def test_discount_totals(self, solved):
        def chk(s):
            _, inv, pct, flat = s
            if flat:
                want = flat if inv.eligible_base > 0 else 0
            else:
                want = inv.eligible_base * pct / 100
            return None if abs(inv.total_discount - want) < 1e-5 else f"discount {inv.total_discount} != {want}"
        assert not _fails(solved, chk)

    @allure.story("Final Price = ROUND(Total - Discount, 0) and Round Off reconciles")
    def test_final_price_rounding(self, solved):
        def chk(s):
            inv = s[1]
            exp = excel_round(inv.total_price - inv.total_discount)
            if inv.final_price != exp:
                return f"final {inv.final_price} != {exp}"
            if abs(inv.round_adjustment - (inv.final_price - (inv.total_price - inv.total_discount))) > TOL:
                return "round-off mismatch"
        assert not _fails(solved, chk)

    @allure.story("Line costs (purchase price x physical qty) add up to the Final Price")
    def test_line_costs_reconcile(self, solved):
        def chk(s):
            inv = s[1]
            costed = [l for l in inv.lines if l.final_line_cost is not None]
            if len(costed) != len(inv.lines):
                return None
            tot = sum(l.final_line_cost for l in costed)
            return None if abs(tot - inv.final_price) < 1e-4 else f"line costs {tot} != final {inv.final_price}"
        assert not _fails(solved, chk)

    @allure.story("Margins: retail margin % = (retail - purchase price/box) / retail")
    def test_margin_formula(self, solved):
        def chk(s):
            for l in s[1].lines:
                if l.purchase_price_box is not None and l.retail_per_box > 0:
                    exp = (l.retail_per_box - l.purchase_price_box) / l.retail_per_box
                    if abs(l.retail_margin_pct - exp) > TOL:
                        return f"margin {l.retail_margin_pct} != {exp}"
        assert not _fails(solved, chk)


@allure.epic("Inventory")
@allure.feature("Opening Stock mode - calculation model")
class TestSupplierBalance:

    @pytest.mark.parametrize("balance,paid,less,due,adv", [(0, 0, 0, 440, 0), (0, 440, 0, 0, 0), (0, 440, 0, 0, 0),
                                                           (100, 200, 0, 340, 0), (0, 600, 0, 0, 160),
                                                           (0, 0, 440, 0, 0), (-100, 0, 0, 340, 0), (-500, 0, 0, 0, 60)])
    @allure.story("Due / Advance = |balance + Final - Paid - Less| on the correct side (application + Default sheet rule)")
    def test_due_advance(self, balance, paid, less, due, adv):
        inv = build([opening_line(5, 10, 100)], supplier_balance=balance, paid=paid, less=less)
        assert inv.final_price == 440
        assert (inv.due, inv.advance) == (due, adv)


@allure.epic("Inventory")
@allure.feature("Opening Stock mode - workbook")
class TestWorkbookRules:
    """Structural checks of the user's workbook (sheet ``Opening Stock``, corrected 2026-10-12): the sheet must keep
    the rules the oracle relies on."""

    @staticmethod
    def _ws():
        return openpyxl.load_workbook(WORKBOOK)["Opening Stock"]

    @allure.story("Supplier Due / Advance (N3 / N4) deduct Less (N5), same as the Default sheet and the application")
    def test_sheet_due_advance_deduct_less(self):
        ws = self._ws()
        assert "N5" in ws["N3"].value and "N5" in ws["N4"].value, (ws["N3"].value, ws["N4"].value)
        inv = build([opening_line(5, 10, 100)], paid=0, less=100)
        assert sheet_due_advance(inv) == (inv.due, inv.advance) == (340, 0)

    @allure.story("Every product row reads the same discount inputs ($J$2:$J$4)")
    def test_every_row_reads_the_discount_inputs(self):
        ws = self._ws()
        bad = [r for r in range(8, 1008) if "$J$3" not in str(ws[f"P{r}"].value) or "$H$" in str(ws[f"P{r}"].value)]
        assert not bad, f"{len(bad)} product rows (first: {bad[:3]}) do not read $J$3"

    @allure.story("I8 is a typed override on purpose (shows that a manual BTP+VAT/Box overrides the % column); rows 9+ are formulas")
    def test_first_row_btp_is_intentional_override(self):
        ws = self._ws()
        assert not str(ws["I8"].value).startswith("=")
        assert str(ws["I9"].value).startswith("=")

    @allure.story("Excel ROUND half away from zero (used for Final Price) differs from Python's banker's rounding at .5")
    @pytest.mark.parametrize("v,expected", [(0.5, 1), (1.5, 2), (2.5, 3), (-0.5, -1)])
    def test_excel_round(self, v, expected):
        assert excel_round(v) == expected
