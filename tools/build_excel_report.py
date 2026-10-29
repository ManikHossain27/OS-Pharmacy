"""Generate PMS_Product_Module_Test_Report.xlsx from the test-case catalog and
the bug-log JSON.

Usage:
    python tools/build_excel_report.py
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "PMS_Product_Module_Test_Report.xlsx"
BUGS = ROOT / "test_data" / "bug_logs.json"

import sys
sys.path.insert(0, str(ROOT))
from tools.test_case_catalog import (  # noqa: E402
    COLUMNS,
    all_rows,
    summary_by_module_feature,
    summary_by_status,
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
STATUS_FILL = {
    "Pass": PatternFill("solid", fgColor="C6EFCE"),
    "Fail": PatternFill("solid", fgColor="FFC7CE"),
    "Blocked": PatternFill("solid", fgColor="FFEB9C"),
    "Not Run": PatternFill("solid", fgColor="D9D9D9"),
}
SEV_FILL = {
    "Critical": PatternFill("solid", fgColor="C00000"),
    "High": PatternFill("solid", fgColor="FF6D6D"),
    "Medium": PatternFill("solid", fgColor="FFD966"),
    "Low": PatternFill("solid", fgColor="A9D08E"),
}
WRAP = Alignment(wrap_text=True, vertical="top")


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}1"


def _autosize(ws, max_width: int = 60) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[letter].width = min(max(12, longest + 2), max_width)


def sheet_test_cases(wb: Workbook) -> None:
    ws = wb.create_sheet("Test Cases")
    ws.append(COLUMNS)
    rows = all_rows()
    status_idx = COLUMNS.index("Status") + 1
    sev_idx = COLUMNS.index("Severity") + 1
    for r in rows:
        ws.append([r.get(col, "") for col in COLUMNS])
        row_no = ws.max_row
        st = r.get("Status", "")
        if st in STATUS_FILL:
            ws.cell(row=row_no, column=status_idx).fill = STATUS_FILL[st]
        sev = r.get("Severity", "")
        if sev in SEV_FILL:
            ws.cell(row=row_no, column=sev_idx).fill = SEV_FILL[sev]
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(row=row_no, column=c).alignment = WRAP
    _style_header(ws, len(COLUMNS))
    _autosize(ws)
    # keep long free-text columns readable
    for name in ("Reproducing Steps", "Expected Result", "Actual Result",
                 "Test Case Description", "Remarks", "Preconditions", "Test Scenario"):
        ws.column_dimensions[get_column_letter(COLUMNS.index(name) + 1)].width = 46


def sheet_summary(wb: Workbook) -> None:
    ws = wb.create_sheet("Test Summary", 0)
    rows = all_rows()
    total = len(rows)
    by_status = summary_by_status()
    passed = by_status.get("Pass", 0)
    failed = by_status.get("Fail", 0)
    blocked = by_status.get("Blocked", 0)
    not_run = by_status.get("Not Run", 0)
    executed = passed + failed

    ws.append(["PMS - Product, Customization, Inventory & Sale Modules - Test Execution Summary"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Application", "https://pms-sql-portal.ems24.co/"])
    ws.append(["Environment", "Staging"])
    ws.append(["Pharmacy under test", "Manik Pharma (id 606 originally, then id 598 from 2026-09-09 onward after 606 began 404-ing - see env notes)"])
    ws.append(["Test dates", "2026-09-08/09 (Products & Customizations) + 2026-09-14/15 (Inventory) + 2026-09-16 (Inventory Return, Stock Adjustment, Customization regression) + 2026-10-27 (Sale Management)"])
    ws.append(["Prepared by", "Senior SQA (automation-assisted)"])
    ws.append(["Env note (2026-09-09)", "BUG-API-001: mid-run API outage (CORS/500/404, pharmacy 606 -> 404) blocked the full pharmacy-scoped re-run; findings from the 04:42-05:19 UTC healthy window are the run of record."])
    ws.append(["Env note (2026-09-14/15)", "The same BUG-API-001 pattern also appears as spurious 'pharmacy not found' skips whenever the reused Playwright auth state (playwright/.auth/state.json) goes stale after ~30-40 min of continuous runs - not an actual outage; deleting that file and rerunning resolves it."])
    ws.append(["Env note (2026-09-16)", "Inventory Return and Stock Adjustment had zero prior coverage and were tested live end-to-end for the first time this session, including completion, cancellation/reversal, and Details-view accuracy against the raw API. Customizations > Product Type create was regression-confirmed still broken (BUG-PTYPE-001, unchanged); Company create and the new Payment Method List sub-page were spot-checked with no regressions found."])
    ws.append(["Env note (2026-10-27)", "Sale Management (Sale Config, Loyalty Config, Customer, POS/Checkout, Hold, Cashier, Sale Record, Sale Return, Dashboard reconciliation) was executed with Playwright MCP (Chromium; Firefox and WebKit for a smoke check) in Manik Pharma; cases are listed under Module 'Sale' (source: docs/SALE_MANAGEMENT.md). Real sales, returns, cancellations, customers, loyalty redemptions (SMS OTP) and cashier approvals were created on staging. Sale Config, Loyalty Config and Invoice Config were changed for the tests and restored to their original values. Cases that need a second user role, a multi-batch product or a PDF parser are Blocked with the reason in Remarks."])
    ws.append([])

    ws.append(["Metric", "Count", "%"])
    hdr = ws.max_row
    data = [
        ("Total Test Cases", total, "100%"),
        ("Passed", passed, f"{passed/total*100:.1f}%" if total else "0%"),
        ("Failed", failed, f"{failed/total*100:.1f}%" if total else "0%"),
        ("Blocked", blocked, f"{blocked/total*100:.1f}%" if total else "0%"),
        ("Not Run", not_run, f"{not_run/total*100:.1f}%" if total else "0%"),
        ("Executed (Pass+Fail)", executed, f"{executed/total*100:.1f}%" if total else "0%"),
        ("Pass % of executed", passed, f"{(passed/executed*100):.1f}%" if executed else "n/a"),
        ("Fail % of executed", failed, f"{(failed/executed*100):.1f}%" if executed else "n/a"),
    ]
    for name, cnt, pct in data:
        ws.append([name, cnt, pct])
    for c in range(1, 4):
        cell = ws.cell(row=hdr, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    ws.append([])
    ws.append(["Module-wise breakdown"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    mf_hdr = ["Module", "Feature", "Total", "Pass", "Fail", "Blocked", "Not Run"]
    ws.append(mf_hdr)
    hrow = ws.max_row
    for c in range(1, len(mf_hdr) + 1):
        ws.cell(row=hrow, column=c).fill = HEADER_FILL
        ws.cell(row=hrow, column=c).font = HEADER_FONT
    for (module, feature), counts in summary_by_module_feature().items():
        t = sum(counts.values())
        ws.append([
            module, feature, t,
            counts.get("Pass", 0), counts.get("Fail", 0),
            counts.get("Blocked", 0), counts.get("Not Run", 0),
        ])

    ws.append([])
    ws.append(["Defect summary by severity"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    bugs = json.loads(BUGS.read_text(encoding="utf-8")) if BUGS.exists() else []
    sev_counts: dict[str, int] = {}
    for b in bugs:
        sev_counts[b["Severity"]] = sev_counts.get(b["Severity"], 0) + 1
    ws.append(["Severity", "Count"])
    hrow = ws.max_row
    for c in (1, 2):
        ws.cell(row=hrow, column=c).fill = HEADER_FILL
        ws.cell(row=hrow, column=c).font = HEADER_FONT
    for sev in ("Critical", "High", "Medium", "Low"):
        ws.append([sev, sev_counts.get(sev, 0)])
    ws.append(["Total defects", len(bugs)])

    _autosize(ws, max_width=48)
    ws.column_dimensions["A"].width = 30


def sheet_bugs(wb: Workbook) -> None:
    ws = wb.create_sheet("Bug Logs")
    bugs = json.loads(BUGS.read_text(encoding="utf-8")) if BUGS.exists() else []
    if not bugs:
        ws.append(["No defects logged."])
        return
    cols = list(bugs[0].keys())
    ws.append(cols)
    sev_idx = cols.index("Severity") + 1 if "Severity" in cols else None
    for b in bugs:
        ws.append([b.get(c, "") for c in cols])
        row_no = ws.max_row
        if sev_idx:
            sev = b.get("Severity", "")
            if sev in SEV_FILL:
                ws.cell(row=row_no, column=sev_idx).fill = SEV_FILL[sev]
        for c in range(1, len(cols) + 1):
            ws.cell(row=row_no, column=c).alignment = WRAP
    _style_header(ws, len(cols))
    _autosize(ws)
    for name in ("Description", "Steps to Reproduce", "Expected Result",
                 "Actual Result", "Remarks", "Preconditions"):
        if name in cols:
            ws.column_dimensions[get_column_letter(cols.index(name) + 1)].width = 50


def main() -> None:
    wb = Workbook()
    wb.remove(wb.active)  # drop default sheet
    sheet_summary(wb)
    sheet_test_cases(wb)
    sheet_bugs(wb)
    wb.save(OUT)
    print(f"Wrote {OUT}  ({len(all_rows())} test cases, "
          f"{len(json.loads(BUGS.read_text(encoding='utf-8'))) if BUGS.exists() else 0} bugs)")


if __name__ == "__main__":
    main()
