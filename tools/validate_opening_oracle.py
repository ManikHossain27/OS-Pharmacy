"""Prove utils.opening_calc matches the workbook's 'Opening Stock' sheet.

Fills a temp copy of the workbook with N random products under several discount scenarios, recalculates it with
pycel and compares every computed column with the oracle (BTP+VAT derived from retail and the Opening Stock %).

The workbook (corrected 2026-10-12) is used as it is: discount columns of every row read $J$2:$J$4 and the
supplier Due / Advance (N3 / N4) deduct Less (N5). Only I (BTP+VAT/Box) is overwritten with the formula, because
I8 is an intentional typed override.

    python tools/validate_opening_oracle.py [N] [seed]
"""
from __future__ import annotations

import random
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl
from pycel import ExcelCompiler

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from utils.opening_calc import build, opening_line, sheet_due_advance  # noqa: E402

XLSX = Path(__file__).resolve().parent.parent / "Knowledge/Inventory Management/PMS_Innventory Calculation (Non Rounded).xlsx"
SHEET = "Opening Stock"
FIRST = 8


def random_lines(rng: random.Random, n: int, *, free_ratio: float, special_ratio: float):
    out = []
    for i in range(n):
        out.append(opening_line(
            rng.choice([1, 2, 3, 5, 7, 10, 25, 50]), rng.choice([1, 2, 3, 5, 6, 10, 12, 14, 20, 24, 30, 60, 100]),
            round(rng.uniform(20, 1200), 2), pct=rng.choice([12, 12, 15, 0, 7.5, 20, 33.3]), exact=True,
            free_qty=rng.choice([1, 5, 10]) if rng.random() < free_ratio else 0,
            wholesale=round(rng.uniform(15, 1100), 2),
            special=round(rng.uniform(1, 40), 2) if rng.random() < special_ratio else 0, name=f"P{i:03d}"))
    return out


def run(n: int, seed: int) -> int:
    rng = random.Random(seed)
    scenarios = [  # (pct, flat, free_ratio, special_ratio, paid, less)
        (0, 0, 0, 0, 0, 0), (10, 0, 0, 0, 0, 0), (7.5, 0, 0.2, 0, 1000, 0), (0, 500, 0.2, 0, 0, 25.5),
        (0, 12345.67, 0.1, 0, 5000, 0), (12, 0, 0.3, 0.1, 0, 100), (0, 800, 0.1, 0.1, 777, 0),
        (100, 0, 0, 0, 0, 0), (0, 10 ** 9, 0, 0, 0, 0), (10, 0, 0.5, 0, 0, 0),
    ]
    bad = checks = 0
    for pct, flat, fr, sr, paid, less in scenarios:
        lines = random_lines(rng, n, free_ratio=fr, special_ratio=sr)
        tmp = Path(tempfile.mkdtemp()) / "wb.xlsx"
        shutil.copy(XLSX, tmp)
        wb = openpyxl.load_workbook(tmp)
        ws = wb[SHEET]
        for r in range(FIRST, FIRST + 400):
            for col in "ABCEFGIKQ":
                ws[f"{col}{r}"].value = None
            ws[f"I{r}"].value = f'=IF(OR(G{r}="",F{r}=""),"",G{r}*(1-F{r}))'
        ws["J2"], ws["J3"], ws["L5"], ws["N5"] = pct / 100, flat, paid, less
        for i, l in enumerate(lines):
            r = FIRST + i
            used_pct = 1 - l.btp_vat_per_box / l.retail_per_box            # the Opening Stock % used for this row
            ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"E{r}"] = l.name, l.box_qty, l.qty_per_box, l.free_qty
            ws[f"F{r}"], ws[f"G{r}"], ws[f"K{r}"], ws[f"Q{r}"] = used_pct, l.retail_per_box, l.wholesale_per_box, l.special_discount
        wb.save(tmp)
        xl = ExcelCompiler(filename=str(tmp))
        inv = build(lines, pct=pct, flat=flat, paid=paid, less=less, less_basis="purchased",
                    supplier_balance=float(xl.evaluate(f"'{SHEET}'!N2")))

        def cell(a):
            return xl.evaluate(f"'{SHEET}'!{a}")

        due, adv = sheet_due_advance(inv)
        for a, exp in (("L3", inv.total_price), ("L2", inv.total_discount), ("L4", inv.final_price),
                       ("J1", inv.round_adjustment), ("J4", inv.eligible_base), ("N3", due), ("N4", adv)):
            got = cell(a)
            checks += 1
            if abs(float(got) - exp) > 1e-6:
                bad += 1
                print(f"MISMATCH pct={pct} flat={flat} {a}: sheet={got} oracle={exp}")
        cols = {"D": "total_qty", "J": "total_btp_vat", "O": "discount_pct", "P": "discount_value",
                "R": "purchase_price_pc", "S": "purchase_price_box", "T": "retail_margin_pct",
                "U": "retail_margin_value", "V": "wholesale_margin_pct", "W": "wholesale_margin_value",
                "Z": "final_line_cost", "X": "less_amount", "Y": "round_adj"}
        for i, l in enumerate(inv.lines):
            r = FIRST + i
            checks += 1
            got_i = cell(f"I{r}")
            if abs(float(got_i) - l.btp_vat_per_box) > 1e-6:
                bad += 1
                print(f"MISMATCH I{r}: sheet={got_i} oracle={l.btp_vat_per_box}")
            for col, attr in cols.items():
                got, exp = cell(f"{col}{r}"), getattr(l, attr)
                checks += 1
                if exp is None:
                    if got not in ("", None):
                        bad += 1
                        print(f"MISMATCH pct={pct} flat={flat} {col}{r}: sheet={got!r} oracle=blank")
                elif isinstance(got, str) or abs(float(got) - exp) > 1e-6:
                    bad += 1
                    print(f"MISMATCH pct={pct} flat={flat} {col}{r}: sheet={got!r} oracle={exp}")
        print(f"scenario pct={pct} flat={flat} free={fr} special={sr}: n={n} final={inv.final_price} "
              f"discount={inv.total_discount:.2f} -> {'OK' if not bad else 'MISMATCHES SO FAR: ' + str(bad)}")
    print(f"\n{checks} cell comparisons, {bad} mismatches")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(run(int(sys.argv[1]) if len(sys.argv) > 1 else 60, int(sys.argv[2]) if len(sys.argv) > 2 else 7))
