"""Prove utils.tpvat_calc matches the workbook's 'TP per vat mode (2)' sheet.

Fills a temp copy of the workbook with N random products (default 300) under a
number of discount / payment / less scenarios, recalculates it with pycel and
compares every computed column with the oracle.

    python tools/validate_tpvat_oracle.py [N] [seed]
"""
from __future__ import annotations

import random
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl
from pycel import ExcelCompiler

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from utils.tpvat_calc import TpVatLine, build, vat_for  # noqa: E402

XLSX = Path(__file__).resolve().parent.parent / "Knowledge/Inventory Management/PMS_Innventory Calculation (Non Rounded).xlsx"
SHEET = "TP per vat mode (2)"
FIRST = 8


def random_lines(rng: random.Random, n: int, *, free_ratio: float, special_ratio: float,
                 manual_vat_ratio: float = 0.3, zero_tp_ratio: float = 0.0) -> list[TpVatLine]:
    out = []
    for i in range(n):
        qpb = rng.choice([1, 2, 3, 5, 6, 10, 12, 14, 20, 24, 30, 60, 100])
        tp = 0.0 if rng.random() < zero_tp_ratio else round(rng.uniform(5, 800), 2)
        vat = round(rng.uniform(0, 120), 2) if rng.random() < manual_vat_ratio else vat_for(tp)
        out.append(TpVatLine(
            name=f"P{i:03d}", box_qty=rng.choice([1, 2, 3, 5, 7, 10, 25, 50]), qty_per_box=qpb,
            tp_per_box=tp, vat_per_box=vat,
            free_qty=rng.choice([1, 5, 10]) if rng.random() < free_ratio else 0,
            retail_per_box=round(rng.uniform(950, 1200), 2), wholesale_per_box=round(rng.uniform(900, 1100), 2),
            special_discount=round(rng.uniform(1, 40), 2) if rng.random() < special_ratio else 0))
    return out


def run(n: int, seed: int) -> int:
    rng = random.Random(seed)
    scenarios = [  # (pct, flat, free_ratio, special_ratio, paid, less, zero_tp_ratio)
        (0, 0, 0, 0, 0, 0, 0), (10, 0, 0, 0, 0, 0, 0), (7.5, 0, 0.2, 0, 1000, 0, 0), (0, 500, 0.2, 0, 0, 25.5, 0),
        (0, 12345.67, 0.1, 0, 5000, 0, 0), (12, 0, 0.3, 0.1, 0, 100, 0), (0, 800, 0.1, 0.1, 777, 0, 0),
        (100, 0, 0, 0, 0, 0, 0), (0, 10 ** 9, 0, 0, 0, 0, 0), (10, 0, 0.5, 0, 0, 0, 0),
        (5, 0, 0.1, 0, 0, 0, 0.1),                                   # zero-TP lines (spec 12 / 41)
    ]
    bad = 0
    checks = 0
    for pct, flat, fr, sr, paid, less, ztp in scenarios:
        lines = random_lines(rng, n, free_ratio=fr, special_ratio=sr, zero_tp_ratio=ztp)
        tmp = Path(tempfile.mkdtemp()) / "wb.xlsx"
        shutil.copy(XLSX, tmp)
        wb = openpyxl.load_workbook(tmp)
        ws = wb[SHEET]
        for r in range(FIRST, FIRST + 400):
            for col in "ABCEFHKMS":
                ws[f"{col}{r}"].value = None
        ws["I1"], ws["I2"], ws["K4"], ws["M5"] = pct / 100, flat, paid, less
        for i, l in enumerate(lines):
            r = FIRST + i
            ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"E{r}"] = l.name, l.box_qty, l.qty_per_box, l.free_qty
            ws[f"F{r}"], ws[f"H{r}"], ws[f"K{r}"], ws[f"M{r}"] = l.tp_per_box, l.vat_per_box, l.retail_per_box, l.wholesale_per_box
            ws[f"S{r}"] = l.special_discount
        wb.save(tmp)
        xl = ExcelCompiler(filename=str(tmp))
        inv = build(lines, pct=pct, flat=flat, paid=paid, less=less, less_rule="sheet",
                    supplier_balance=float(xl.evaluate(f"'{SHEET}'!M2")))

        def cell(a):
            return xl.evaluate(f"'{SHEET}'!{a}")

        for a, exp in (("G2", inv.total_price), ("G3", inv.total_tp), ("G4", inv.total_vat), ("I3", inv.eligible_tp),
                       ("I4", inv.total_discount), ("K1", inv.k1), ("K2", inv.final_price),
                       ("K3", inv.round_adjustment), ("M3", inv.due), ("M4", inv.advance)):
            got = cell(a); checks += 1
            if abs(float(got) - exp) > 1e-6:
                bad += 1; print(f"MISMATCH pct={pct} flat={flat} {a}: sheet={got} oracle={exp}")
        cols = {"D": "total_qty", "J": "total_tp", "AC": "total_price", "Q": "discount_pct", "R": "discount_value",
                "T": "purchase_price_pc", "U": "purchase_price_box", "V": "retail_margin_pct",
                "W": "retail_margin_value", "X": "wholesale_margin_pct", "Y": "wholesale_margin_value",
                "Z": "less_amount", "AA": "round_adj", "AB": "final_line_cost"}
        for i, l in enumerate(inv.lines):
            for col, attr in cols.items():
                got, exp = cell(f"{col}{FIRST + i}"), getattr(l, attr)
                checks += 1
                if col == "T" and l.sheet_error:                    # the sheet itself errors (#VALUE! / #DIV/0!)
                    if not (isinstance(got, str) and got.startswith("#")) and "Error" not in type(got).__name__ \
                            and not str(got).startswith("#"):
                        bad += 1; print(f"MISMATCH {col}{FIRST+i}: sheet={got!r} oracle={l.sheet_error}")
                    continue
                if exp is None:
                    if got not in ("", None) and not str(got).startswith("#"):
                        bad += 1; print(f"MISMATCH pct={pct} flat={flat} {col}{FIRST+i}: sheet={got!r} oracle=blank")
                elif isinstance(got, str) or abs(float(got) - exp) > 1e-6:
                    if l.sheet_error and col in ("U", "V", "W", "X", "Y", "AB"):
                        continue                                        # downstream of the sheet's own error
                    bad += 1; print(f"MISMATCH pct={pct} flat={flat} {col}{FIRST+i}: sheet={got!r} oracle={exp}")
        print(f"scenario pct={pct} flat={flat} free={fr} special={sr} zeroTP={ztp}: n={n} final={inv.final_price} "
              f"discount={inv.total_discount:.2f} -> {'OK' if not bad else 'MISMATCHES SO FAR: ' + str(bad)}")
    print(f"\n{checks} cell comparisons, {bad} mismatches")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(run(int(sys.argv[1]) if len(sys.argv) > 1 else 300, int(sys.argv[2]) if len(sys.argv) > 2 else 7))
